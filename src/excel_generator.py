from pathlib import Path

import pandas as pd
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.utils import get_column_letter


# ============================================================
# OUTPUT CONFIGURATION
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent

OUTPUT_DIR = PROJECT_ROOT / "output"

DEFAULT_OUTPUT_FILE = (
    OUTPUT_DIR
    / "enterprise_automation_report.xlsx"
)


# ============================================================
# SUMMARY
# ============================================================

def build_summary(dataframe):
    """
    Build high-level processing metrics for the SUMMARY sheet.
    """

    total_records = len(dataframe)

    classification_counts = (
        dataframe["classification"]
        .value_counts()
    )

    validation_counts = (
        dataframe["validation_status"]
        .value_counts()
    )

    summary_rows = [
        {
            "metric": "Total Records",
            "value": total_records,
        },
        {
            "metric": "Standard",
            "value": classification_counts.get(
                "STANDARD",
                0,
            ),
        },
        {
            "metric": "High Priority",
            "value": classification_counts.get(
                "HIGH PRIORITY",
                0,
            ),
        },
        {
            "metric": "Review",
            "value": classification_counts.get(
                "REVIEW",
                0,
            ),
        },
        {
            "metric": "Exceptions",
            "value": classification_counts.get(
                "EXCEPTION",
                0,
            ),
        },
        {
            "metric": "Validation Passed",
            "value": validation_counts.get(
                "PASSED",
                0,
            ),
        },
        {
            "metric": "Validation Warnings",
            "value": validation_counts.get(
                "WARNING",
                0,
            ),
        },
        {
            "metric": "Validation Failed",
            "value": validation_counts.get(
                "FAILED",
                0,
            ),
        },
    ]

    return pd.DataFrame(
        summary_rows
    )


# ============================================================
# OUTPUT VIEWS
# ============================================================

def build_priority_view(dataframe):
    """
    Return records classified as HIGH PRIORITY.
    """

    return (
        dataframe[
            dataframe["classification"]
            == "HIGH PRIORITY"
        ]
        .copy()
    )


def build_exception_view(dataframe):
    """
    Return records requiring business or data-quality review.
    """

    business_issue = (
        dataframe["classification"]
        .isin(
            [
                "REVIEW",
                "EXCEPTION",
            ]
        )
    )

    validation_issue = (
        dataframe["validation_status"]
        != "PASSED"
    )

    return (
        dataframe[
            business_issue
            | validation_issue
        ]
        .copy()
    )


# ============================================================
# EXCEL FORMATTING
# ============================================================

def format_worksheet(worksheet):
    """
    Apply basic professional formatting to an Excel worksheet.
    """

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in worksheet[1]:

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    for column_cells in worksheet.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0

        for cell in column_cells:

            cell_value = cell.value

            if cell_value is None:
                continue

            max_length = max(
                max_length,
                len(str(cell_value)),
            )

        adjusted_width = min(
            max(max_length + 2, 12),
            45,
        )

        worksheet.column_dimensions[
            column_letter
        ].width = adjusted_width


def format_summary_sheet(worksheet):
    """
    Apply additional formatting to the SUMMARY worksheet.
    """

    metric_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    for row in range(
        2,
        worksheet.max_row + 1,
    ):

        worksheet.cell(
            row=row,
            column=1,
        ).fill = metric_fill

        worksheet.cell(
            row=row,
            column=1,
        ).font = Font(
            bold=True
        )


# ============================================================
# EXCEL GENERATION
# ============================================================

def generate_excel_report(
    dataframe,
    output_file=None,
):
    """
    Generate the final Excel business deliverable.

    Workbook sheets:
        SUMMARY
        MASTER
        PRIORITIES
        EXCEPTIONS
    """

    if output_file is None:
        output_file = DEFAULT_OUTPUT_FILE

    output_file = Path(
        output_file
    )

    output_file.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    summary = build_summary(
        dataframe
    )

    priorities = build_priority_view(
        dataframe
    )

    exceptions = build_exception_view(
        dataframe
    )

    with pd.ExcelWriter(
        output_file,
        engine="openpyxl",
    ) as writer:

        summary.to_excel(
            writer,
            sheet_name="SUMMARY",
            index=False,
        )

        dataframe.to_excel(
            writer,
            sheet_name="MASTER",
            index=False,
        )

        priorities.to_excel(
            writer,
            sheet_name="PRIORITIES",
            index=False,
        )

        exceptions.to_excel(
            writer,
            sheet_name="EXCEPTIONS",
            index=False,
        )

        workbook = writer.book

        for worksheet in workbook.worksheets:

            format_worksheet(
                worksheet
            )

        format_summary_sheet(
            workbook["SUMMARY"]
        )

    return output_file
