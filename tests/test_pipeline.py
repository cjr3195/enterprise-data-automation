from pathlib import Path

from openpyxl import load_workbook

from src.extract import load_source_data
from src.transform import transform_source_data
from src.business_rules import apply_business_rules
from src.validation import validate_dataset
from src.excel_generator import generate_excel_report


def test_end_to_end_pipeline(tmp_path):
    """
    Validate the complete automation pipeline from
    synthetic source data through Excel generation.
    """

    # --------------------------------------------------------
    # EXTRACT
    # --------------------------------------------------------

    source_data = load_source_data()

    assert "customers" in source_data
    assert "contracts" in source_data
    assert "priority_rules" in source_data

    assert len(source_data["customers"]) > 0
    assert len(source_data["contracts"]) > 0
    assert len(source_data["priority_rules"]) > 0

    # --------------------------------------------------------
    # TRANSFORM
    # --------------------------------------------------------

    transformed_data = transform_source_data(
        source_data
    )

    assert len(transformed_data) > 0

    assert "customer_id" in transformed_data.columns
    assert "contract_id" in transformed_data.columns
    assert "priority_level" in transformed_data.columns

    # --------------------------------------------------------
    # BUSINESS RULES
    # --------------------------------------------------------

    evaluated_data = apply_business_rules(
        transformed_data
    )

    assert "classification" in evaluated_data.columns
    assert "decision_reason" in evaluated_data.columns
    assert "rule_applied" in evaluated_data.columns

    assert (
        evaluated_data["classification"]
        .notna()
        .all()
    )

    # --------------------------------------------------------
    # VALIDATION
    # --------------------------------------------------------

    validated_data = validate_dataset(
        evaluated_data
    )

    assert "validation_status" in validated_data.columns
    assert "validation_issues" in validated_data.columns

    assert (
        validated_data["validation_status"]
        .notna()
        .all()
    )

    # --------------------------------------------------------
    # EXCEL GENERATION
    # --------------------------------------------------------

    output_file = (
        tmp_path
        / "enterprise_automation_report.xlsx"
    )

    generated_file = generate_excel_report(
        validated_data,
        output_file=output_file,
    )

    generated_file = Path(
        generated_file
    )

    assert generated_file.exists()

    # --------------------------------------------------------
    # EXCEL VALIDATION
    # --------------------------------------------------------

    workbook = load_workbook(
        generated_file,
        read_only=True,
    )

    expected_sheets = {
        "SUMMARY",
        "MASTER",
        "PRIORITIES",
        "EXCEPTIONS",
    }

    assert expected_sheets.issubset(
        set(workbook.sheetnames)
    )

    master_sheet = workbook["MASTER"]

    # Header + processed records
    expected_rows = (
        len(validated_data)
        + 1
    )

    assert (
        master_sheet.max_row
        == expected_rows
    )

    workbook.close()
